"""Verify pinned Texas annual listings and annotate matching current notices.

These listings name a WARN notice day and an anticipated layoff day.  The
agency receipt field is retained as source text and never substitutes for a
notice day.  This module does not infer event identity from names or dates.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from warnlive.store.dedupe import append_repair_version

HEADERS = ("NOTICE_DATE", "JOB_SITE_NAME", "COUNTY_NAME", "WDA_NAME",
           "TOTAL_LAYOFF_NUMBER", "LayOff_Date", "WFDD_RECEIVED_DATE", "CITY_NAME")
YEARS = range(2020, 2027)
RULE = "tx_annual_workbook_dates_v1"


def _source_cell(value: object) -> dict:
    if isinstance(value, (date, datetime)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, bool):
        raise ValueError("Texas workbook has a boolean source cell")
    if isinstance(value, (str, int, float)):
        return {"type": type(value).__name__, "value": value}
    if value is None:
        return {"type": "blank", "value": None}
    raise ValueError(f"unsupported Texas source cell: {type(value).__name__}")


def _day(value: object) -> str | None:
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ValueError("Texas workbook date has a time component")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    raise ValueError(f"Texas workbook date is not an Excel date: {value!r}")


def _fields(raw: dict) -> tuple[str, ...]:
    extra = set(raw) - set(HEADERS)
    if (set(HEADERS) - set(raw) or extra - {None, "_restkey"}
            or any(cell not in (None, "") for key, value in raw.items() if key in extra
                   for cell in (value if isinstance(value, list) else [value]))):
        raise ValueError("Texas row does not have exactly eight named fields")
    return tuple(raw[name] or "" for name in HEADERS)


def read_artifacts(directory: Path, raw_csv: Path | None = None) -> list[dict]:
    """Fail closed on workbook bytes, sheet layout, cells, and raw correspondence."""
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if manifest.get("format") != "warn-tx-annual-artifacts-v1":
        raise ValueError("unsupported Texas source manifest")
    page = directory / "source.html"
    if page.is_symlink() or not page.is_file():
        raise ValueError("Texas source-page snapshot is missing")
    page_bytes = page.read_bytes()
    if hashlib.sha256(page_bytes).hexdigest() != manifest.get("source_page_snapshot_sha256"):
        raise ValueError("Texas source-page snapshot drift")
    html = page_bytes.decode("utf-8")
    artifacts = manifest.get("artifacts")
    if (not isinstance(artifacts, list) or len(artifacts) != len(YEARS)
            or {item.get("path") for item in artifacts} != {f"{year}.xlsx" for year in YEARS}):
        raise ValueError("Texas manifest must list seven annual workbooks")
    rows: list[dict] = []
    for item in sorted(artifacts, key=lambda entry: entry["path"]):
        name = item["path"]
        path = directory / name
        if path.is_symlink() or not path.is_file():
            raise ValueError(f"invalid Texas workbook: {path}")
        content = path.read_bytes()
        digest = hashlib.sha256(content).hexdigest()
        if digest != item.get("sha256") or len(content) != item.get("bytes"):
            raise ValueError(f"Texas workbook checksum mismatch: {name}")
        if (item.get("headers") != list(HEADERS)
                or item.get("year") != int(name[:4])
                or not isinstance(item.get("source_url"), str)
                or not item["source_url"].startswith("https://www.twc.texas.gov/")):
            raise ValueError(f"Texas workbook manifest schema mismatch: {name}")
        url_path = item["source_url"].removeprefix("https://www.twc.texas.gov")
        if (not re.search(rf'href="{re.escape(url_path)}"', html)
                or not url_path.endswith(f"warn-act-listings-{item['year']}-twc.xlsx")):
            raise ValueError(f"Texas workbook source URL is absent from page: {name}")
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if book.sheetnames != [item.get("sheet")]:
                raise ValueError(f"Texas workbook sheet layout changed: {name}")
            sheet = book.active
            if sheet.max_column != item.get("columns") or sheet.max_row != item.get("max_row"):
                raise ValueError(f"Texas workbook dimensions changed: {name}")
            header = list(next(sheet.iter_rows(max_row=1, values_only=True)))
            if tuple(header[:8]) != HEADERS or any(cell is not None for cell in header[8:]):
                raise ValueError(f"Texas workbook header changed: {name}")
            count = 0
            trailing_blank = False
            for ordinal, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if all(cell is None for cell in values):
                    trailing_blank = True
                    continue
                if trailing_blank or any(cell not in (None, "") for cell in values[8:]):
                    raise ValueError(f"Texas workbook unexpected cell at {name}:{ordinal}")
                cells = values[:8]
                if cells[0] is None or cells[1] is None:
                    raise ValueError(f"Texas workbook incomplete row: {name}:{ordinal}")
                selected = tuple(str(cell) if cell is not None else "" for cell in cells)
                notice, effective = _day(cells[0]), _day(cells[5])
                _day(cells[6])
                if notice is None:
                    raise ValueError(f"Texas workbook notice date missing: {name}:{ordinal}")
                rows.append({
                    "source_artifact": f"agency/tx/{name}",
                    "source_artifact_sha256": digest,
                    "source_url": item["source_url"],
                    "source_sheet": sheet.title,
                    "source_row": ordinal,
                    "source_fields": dict(zip(HEADERS, selected, strict=True)),
                    "raw_cells": [_source_cell(cell) for cell in cells],
                    "notice_date": notice, "effective_date": effective,
                })
                count += 1
            if count != item.get("data_rows"):
                raise ValueError(f"Texas workbook row count changed: {name}")
        finally:
            book.close()
    if raw_csv is not None:
        raw_csv = Path(raw_csv)
        content = raw_csv.read_bytes()
        if (len(content) != manifest.get("raw_csv_bytes")
                or hashlib.sha256(content).hexdigest() != manifest.get("raw_csv_sha256")):
            raise ValueError("Texas pinned raw CSV drift")
        with raw_csv.open(newline="") as stream:
            reader = csv.DictReader(stream)
            if reader.fieldnames != list(HEADERS):
                raise ValueError("Texas pinned raw CSV header drift")
            raw_counts = Counter(_fields(row) for row in reader)
        source_counts = Counter(_fields(row["source_fields"]) for row in rows)
        if any(raw_counts[key] != count for key, count in source_counts.items()):
            raise ValueError("Texas workbook/raw eight-field correspondence drift")
    return rows


def apply_date_evidence(
    conn: sqlite3.Connection, directory: Path, raw_csv: Path, observed_at: str,
    exceptions: list[dict],
) -> dict:
    """Annotate a unique current raw match; hold mismatches in the ledger."""
    if conn.row_factory is not sqlite3.Row:
        raise ValueError("Texas candidate connection must use sqlite3.Row")
    official = read_artifacts(directory, raw_csv)
    by_fields: dict[tuple[str, ...], list[dict]] = defaultdict(list)
    for source in official:
        by_fields[_fields(source["source_fields"])].append(source)
    candidate: dict[tuple[str, ...], list[sqlite3.Row]] = defaultdict(list)
    for row in conn.execute(
        "SELECT n.*,v.fields_json FROM notices n JOIN notice_versions v "
        "ON v.notice_id=n.id AND v.version=n.current_version "
        "WHERE n.state='TX' AND n.source_identity IS NULL"
    ):
        version = json.loads(row["fields_json"])
        candidate[_fields(json.loads(version["raw_extra"]))].append(row)
    report = {"source_occurrences": len(official), "unique_source_rows": len(by_fields),
              "matched_current_notices": 0, "held_source_rows": 0,
              "reported_notice_dates": 0, "reported_effective_dates": 0,
              "held_corrected_effective_dates": 0, "negative_intervals": 0,
              "raw_anomalous_year_rows_outside_annual_corpus": 0,
              "repeated_source_occurrences": sum(len(v) - 1 for v in by_fields.values())}
    prepared = []
    for fields, sources in sorted(by_fields.items()):
        matches = candidate.get(fields, [])
        first = sources[0]
        pointer = [{"artifact": r["source_artifact"],
                    "sha256": r["source_artifact_sha256"], "sheet": r["source_sheet"],
                    "row": r["source_row"]} for r in sources]
        reason = None
        if len(matches) != 1:
            reason = "no_current_raw_match" if not matches else "ambiguous_current_raw_match"
        else:
            row = matches[0]
            version_raw_fields = {
                _fields(json.loads(json.loads(version[0])["raw_extra"]))
                for version in conn.execute(
                    "SELECT fields_json FROM notice_versions WHERE notice_id=?",
                    (row["id"],),
                )
            }
            if len(version_raw_fields) > 1:
                reason = "same_key_version_raw_conflict"
            elif row["notice_date"] != first["notice_date"]:
                reason = "canonical_notice_date_mismatch"
            elif row["notice_date_precision"] or row["notice_date_basis"]:
                reason = "existing_notice_date_evidence"
        if reason:
            report["held_source_rows"] += len(sources)
            exceptions.append({"origin": "agency/tx", "state": "TX", "reason": reason,
                               "source_rows": pointer, "source_fields": first["source_fields"],
                               "candidate_keys": [r["dedupe_key"] for r in matches]})
            continue
        anomalous_year = (first["effective_date"] is not None
                          and first["effective_date"][:4] in {"1930", "2027"})
        effective_ok = (first["effective_date"] is not None and not anomalous_year
                        and row["effective_date"] == first["effective_date"]
                        and row["effective_date_end"] is None
                        and not row["effective_date_precision"]
                        and not row["effective_date_basis"])
        if first["effective_date"] is not None and not effective_ok:
            report["held_corrected_effective_dates"] += 1
            exceptions.append({"origin": "agency/tx", "state": "TX",
                               "reason": ("anomalous_source_effective_year_requires_review" if anomalous_year
                                          else "effective_date_mismatch_or_existing_evidence"),
                               "source_rows": pointer, "source_fields": first["source_fields"],
                               "candidate_keys": [row["dedupe_key"]],
                               "candidate_effective_date": row["effective_date"]})
        if first["effective_date"] and first["notice_date"] > first["effective_date"]:
            report["negative_intervals"] += 1
        prepared.append((row, first, pointer, effective_ok))
    conn.execute("SAVEPOINT tx_annual_evidence")
    try:
        for row, source, pointer, effective_ok in prepared:
            details = json.loads(row["source_details"] or "{}")
            if (not isinstance(details, dict) or "tx_annual_workbook" in details
                    or details.get("date_evidence_rule") not in (None, RULE)):
                raise ValueError(f"Texas source details conflict: {row['dedupe_key']}")
            details["date_evidence_rule"] = RULE
            details["tx_annual_workbook"] = {
                "date_evidence_rule": RULE,
                "source_rows": pointer,
                "source_fields": source["source_fields"],
                "raw_cells": source["raw_cells"],
                "notice_date_role": "listed_warn_notice_date",
                "effective_date_role": "anticipated_layoff_date",
                "receipt_date_role": "agency_received_date_not_notice_date",
                "effective_date_review": ("source_day_matches_canonical" if effective_ok
                                          else "held_source_or_canonical_mismatch"),
                "negative_notice_to_layoff_interval": bool(
                    source["effective_date"] and source["notice_date"] > source["effective_date"]
                ),
            }
            conn.execute(
                "UPDATE notices SET source_details=?, notice_date_precision='day', "
                "notice_date_basis='reported', effective_date_precision=?, "
                "effective_date_basis=? WHERE id=?",
                (json.dumps(details, sort_keys=True, ensure_ascii=False),
                 "day" if effective_ok else row["effective_date_precision"],
                 "reported" if effective_ok else row["effective_date_basis"], row["id"]),
            )
            if not append_repair_version(conn, row["id"], observed_at):
                raise ValueError(f"Texas evidence produced no version: {row['dedupe_key']}")
            report["matched_current_notices"] += 1
            report["reported_notice_dates"] += 1
            report["reported_effective_dates"] += int(effective_ok)
        conn.execute("RELEASE SAVEPOINT tx_annual_evidence")
    except Exception:
        conn.execute("ROLLBACK TO SAVEPOINT tx_annual_evidence")
        conn.execute("RELEASE SAVEPOINT tx_annual_evidence")
        raise
    # A pre-2020 raw row with a 2027 layoff cell is outside this annual
    # source corpus. Retain an explicit review pointer; do not borrow the
    # later workbook's date authority for the historical record.
    with Path(raw_csv).open(newline="") as stream:
        reader = csv.DictReader(stream)
        for ordinal, raw in enumerate(reader, start=2):
            fields = _fields(raw)
            if fields in by_fields or not fields[5].startswith(("1930-", "2027-")):
                continue
            matches = candidate.get(fields, [])
            exceptions.append({
                "origin": "raw/tx.csv", "state": "TX",
                "reason": "anomalous_raw_effective_year_outside_annual_corpus",
                "raw_row": ordinal, "source_fields": dict(zip(HEADERS, fields, strict=True)),
                "candidate_keys": [row["dedupe_key"] for row in matches],
            })
            report["raw_anomalous_year_rows_outside_annual_corpus"] += 1
    return report
