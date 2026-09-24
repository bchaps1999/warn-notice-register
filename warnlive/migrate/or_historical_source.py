"""Agency-sent Oregon historical WARN list, preserved by Big Local News."""

from __future__ import annotations

import gzip
import hashlib
import io
import json
import tarfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from warnlive.migrate.or_source import HEADERS
from warnlive.migrate.source_bundle import _entry, verify
from warnlive.normalize.engine import _record_hash

ARTIFACT = "or_warnlist_july_2021.xlsx"
PREFIX = "agency/or_historical"
# These pinned rows put only a street/city in Company Name and have no location.
# The employer cannot be recovered from this workbook alone.
UNRESOLVED_EMPLOYER_IDS = {"0826", "0810", "0809", "0742", "0727", "0708"}


def _json(value: object) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, default=str,
                      separators=(",", ":"))


def _day(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError(f"Oregon historical date is not a day: {value!r}") from exc
    if not isinstance(value, datetime) or value.time() != datetime.min.time():
        raise ValueError(f"Oregon historical date is not a day: {value!r}")
    # Excel's zero-date sentinel is an empty action date, not a 19th-century layoff.
    if value.date().isoformat() == "1899-12-29":
        return None
    return value.date().isoformat()


def read_artifacts(directory: Path) -> tuple[list[dict], dict]:
    """Check exact workbook bytes/layout and preserve every physical data row."""
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    path = directory / ARTIFACT
    if (manifest.get("format") != "agency-sent-historical-workbook-v1"
            or manifest.get("file") != ARTIFACT
            or manifest.get("sheet") != "WARNList"
            or manifest.get("headers") != list(HEADERS)
            or manifest.get("data_rows") != 1082
            or path.is_symlink()):
        raise ValueError("unsupported Oregon historical manifest")
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    if len(content) != manifest.get("bytes") or digest != manifest.get("sha256"):
        raise ValueError("Oregon historical workbook checksum mismatch")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if book.sheetnames != ["WARNList"]:
            raise ValueError("Oregon historical workbook sheet changed")
        values = list(book.active.values)
    finally:
        book.close()
    if (len(values) != 1085 or tuple(values[2]) != HEADERS
            or values[1][1] != manifest.get("banner")):
        raise ValueError("Oregon historical workbook layout changed")
    rows = []
    for ordinal, cells in enumerate(values[3:], start=4):
        if len(cells) != len(HEADERS):
            raise ValueError(f"Oregon historical row width changed: {ordinal}")
        raw = {key: value.isoformat() if isinstance(value, datetime) else value
               for key, value in zip(HEADERS, cells, strict=True)}
        rows.append({"source_row": ordinal, "raw": raw,
                     "source_row_sha256": hashlib.sha256(_json(raw).encode()).hexdigest()})
    return rows, manifest


def project(directory: Path, existing_ids: set[str] | None = None) -> tuple[list[dict], list[dict], dict]:
    """Admit complete unique WARN numbers absent from both newer captures."""
    rows, manifest = read_artifacts(directory)
    existing = existing_ids or set()
    counts = Counter(str(row["raw"]["WARN#"] or "").strip() for row in rows)
    records, held = [], []
    for row in rows:
        raw = row["raw"]
        ident = str(raw["WARN#"] or "").strip()
        employer = str(raw["Company Name"] or "").strip()
        try:
            received = _day(raw["Received Date"])
            effective = _day(raw["Layoff Date"])
        except ValueError:
            received = effective = None
        workers = raw["Laid Off"]
        complete = (ident.isdigit() and employer and received and effective
                    and isinstance(workers, (int, float)) and not isinstance(workers, bool)
                    and workers > 0 and int(workers) == workers)
        reason = ("missing_warn_number" if not ident.isdigit() else
                  "newer_agency_capture_overlap" if ident in existing else
                  "multi_site_or_phase_identity_unresolved" if counts[ident] != 1 else
                  "unresolved_employer_in_source" if ident in UNRESOLVED_EMPLOYER_IDS else
                  "incomplete_historical_row" if not complete else None)
        pointer = f"{PREFIX}/{ARTIFACT}:sha256:{manifest['sha256']}:row:{row['source_row']}"
        if reason:
            held.append({"origin": f"{PREFIX}/{ARTIFACT}", "state": "OR",
                         "reason": reason, "source_row": pointer,
                         "source_row_sha256": row["source_row_sha256"],
                         "source_notice_id": ident or None,
                         "source_url": manifest["source_url"],
                         "notice_year": received[:4] if received else None,
                         "raw_extra": _json(raw)})
            continue
        details = {"origin": f"{PREFIX}/{ARTIFACT}", "source_row": pointer,
                   "source_row_sha256": row["source_row_sha256"],
                   "source_workbook_sha256": manifest["sha256"],
                   "provenance_url": manifest["provenance_url"],
                   "identity_basis": "singleton_agency_warn_number",
                   "agency_received_date": received,
                   "date_roles": {"Received Date": "agency_receipt",
                                  "Layoff Date": "reported_action"},
                   "raw_cells": raw}
        record = {"state": "OR", "employer_name": employer,
                  "location": str(raw["Location"] or "").strip() or None,
                  "notice_date": None, "effective_date": effective,
                  "effective_date_precision": "day", "effective_date_basis": "reported",
                  "employees_affected": int(workers), "layoff_type": "unknown",
                  "is_temporary": None, "is_amendment": 0,
                  "source_url": manifest["source_url"], "source_notice_id": ident,
                  "source_identity": f"OR:agency:{ident}",
                  "source_details": _json(details), "raw_extra": _json(raw),
                  "dedupe_key": hashlib.sha1(f"OR|agency|{ident}".encode()).hexdigest()}
        record["raw_record_hash"] = _record_hash(record)
        records.append(record)
    if len(records) + len(held) != len(rows):
        raise ValueError("Oregon historical source accounting mismatch")
    return records, held, {"source_rows": len(rows), "admitted": len(records),
                           "held": len(held),
                           "hold_reasons": dict(Counter(item["reason"] for item in held))}


def build(base_bundle: Path, artifacts: Path, out_bundle: Path) -> dict:
    """Add the pinned original workbook and provenance without overwriting sources."""
    base_bundle, artifacts, out_bundle = map(Path, (base_bundle, artifacts, out_bundle))
    if out_bundle.exists():
        raise FileExistsError(out_bundle)
    base = verify(base_bundle)
    if base.get("admission_inputs") != "agency-only-v1":
        raise ValueError("Oregon historical overlay requires an agency-only base")
    read_artifacts(artifacts)
    with tarfile.open(base_bundle, "r:gz") as archive:
        files = {member.name: archive.extractfile(member).read()
                 for member in archive.getmembers() if member.name != "manifest.json"}
    if any(name.startswith(PREFIX + "/") for name in files):
        raise ValueError("base bundle already contains Oregon historical source")
    for name in ("manifest.json", ARTIFACT):
        files[f"{PREFIX}/{name}"] = (artifacts / name).read_bytes()
    manifest = {**base, "files": [
        {"path": name, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        for name, content in sorted(files.items())]}
    out_bundle.parent.mkdir(parents=True, exist_ok=True)
    try:
        with out_bundle.open("xb") as raw, gzip.GzipFile(fileobj=raw, mode="wb", filename="", mtime=0) as zipped, tarfile.open(fileobj=zipped, mode="w") as archive:
            header = (json.dumps(manifest, sort_keys=True, indent=2) + "\n").encode()
            archive.addfile(_entry("manifest.json", header), io.BytesIO(header))
            for name, content in sorted(files.items()):
                archive.addfile(_entry(name, content), io.BytesIO(content))
        verify(out_bundle)
    except BaseException:
        out_bundle.unlink(missing_ok=True)
        raise
    return manifest
