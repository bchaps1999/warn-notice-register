"""Pinned TWC historical workbook supplied to Stanford Big Local News in 2021.

The mirrored workbook is agency-origin evidence.  Every worksheet row is
accounted for, including rows withheld from automatic canonical admission.
"""

from __future__ import annotations

import gzip
import hashlib
import io
import json
import tarfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

from warnlive.migrate.source_bundle import _entry, verify
from warnlive.normalize.engine import _record_hash

HEADERS = (
    "JOB_SITE_NAME", "CITY_NAME", "COUNTY_NAME", "STATE_WDA_ID",
    "LAYOFF_REASON_DESCRIPTION", "WDA_NAME", "TOTAL_LAYOFF_NUMBER",
    "LayOff_Date", "NOTICE_DATE", "INITIAL_VISIT_DATE", "SSA_NAME",
    "NOTICE_ID", "WFDD_RECEIVED_DATE", "DISTRICT_NUMBER",
    "VISIT_STATUS_CODE_DESCRIPTION", "MANUF", "WARN_COMPLIED",
    "Temporary_Layoff_Flag", "Recall_Date", "TempLay01", "TempLayoffNumber",
)
EXPECTED_ROWS = 5097
SOURCE_NAME = "agency/tx_historical/tx_historical.xlsx"
EXPECTED_SHA256 = "0ee280d92a7091d281b6bf8e43a84f38c27bebda3a493545d1cc22a9a799cfc8"


def _json(value: object) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _raw(value: object) -> object:
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _date(value: object, field: str, row_number: int, optional: bool = False) -> str | None:
    if value is None and optional:
        return None
    if not isinstance(value, datetime) or value.time().isoformat() != "00:00:00":
        raise ValueError(f"Texas workbook invalid {field} at row {row_number}")
    return value.date().isoformat()


def read_artifacts(directory: Path) -> tuple[list[dict], dict]:
    """Verify workbook bytes and the worksheet identity, layout, and all rows."""
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if (manifest.get("format") != "agency-sent-historical-workbook-v1"
            or manifest.get("file") != "tx_historical.xlsx"
            or manifest.get("sha256") != EXPECTED_SHA256):
        raise ValueError("unsupported Texas historical workbook manifest")
    path = directory / "tx_historical.xlsx"
    if path.is_symlink():
        raise ValueError("Texas workbook symlink is not allowed")
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    if digest != EXPECTED_SHA256 or len(content) != manifest.get("bytes"):
        raise ValueError("Texas historical workbook checksum mismatch")
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("Texas workbook worksheet layout changed")
        sheet = workbook["Sheet1"]
        if sheet.max_row != EXPECTED_ROWS + 1 or sheet.max_column != len(HEADERS):
            raise ValueError("Texas workbook dimensions changed")
        iterator = sheet.values
        if tuple(next(iterator)) != HEADERS:
            raise ValueError("Texas workbook headers changed")
        result, seen_ids = [], set()
        for row_number, cells in enumerate(iterator, start=2):
            if len(cells) != len(HEADERS):
                raise ValueError(f"Texas workbook row layout changed: {row_number}")
            fields = dict(zip(HEADERS, cells))
            ident = fields["NOTICE_ID"]
            if not isinstance(ident, int) or isinstance(ident, bool) or ident <= 0 or ident in seen_ids:
                raise ValueError(f"Texas workbook NOTICE_ID invalid or repeated: {row_number}")
            seen_ids.add(ident)
            if fields["WARN_COMPLIED"] is not True:
                raise ValueError(f"Texas workbook WARN_COMPLIED changed: {row_number}")
            notice = _date(fields["NOTICE_DATE"], "NOTICE_DATE", row_number)
            if not 1999 <= int(notice[:4]) <= 2019:
                raise ValueError(f"Texas workbook notice year changed: {row_number}")
            received = _date(fields["WFDD_RECEIVED_DATE"], "WFDD_RECEIVED_DATE", row_number)
            action = _date(fields["LayOff_Date"], "LayOff_Date", row_number, optional=True)
            workers = fields["TOTAL_LAYOFF_NUMBER"]
            if workers is not None and (not isinstance(workers, int) or isinstance(workers, bool) or workers < 0):
                raise ValueError(f"Texas workbook worker count changed: {row_number}")
            raw_fields = {key: _raw(value) for key, value in fields.items()}
            row_hash = hashlib.sha256(_json(raw_fields).encode()).hexdigest()
            result.append({"row_number": row_number, "source_notice_id": str(ident),
                           "notice_date": notice, "effective_date": action,
                           "agency_received_date": received, "workers": workers,
                           "fields": raw_fields, "row_sha256": row_hash,
                           "source_row": f"{SOURCE_NAME}:sha256:{digest}:Sheet1:row:{row_number}:NOTICE_ID:{ident}",
                           "artifact_sha256": digest})
        if len(result) != EXPECTED_ROWS:
            raise ValueError("Texas workbook row count changed")
        return result, manifest
    finally:
        workbook.close()


def project(directory: Path, existing_ids: set[str] | None = None,
            existing_events: set[tuple[str, str, str, str]] | None = None
            ) -> tuple[list[dict], list[dict], dict]:
    """Project numbered notices, holding anomalous dates and existing events.

    ``existing_events`` contains normalized (employer, notice date, city,
    county) tuples from previously admitted Texas records.  Caller performs
    reconciliation after its existing agency ingestion.
    """
    rows, manifest = read_artifacts(directory)
    existing_ids = {str(value).removeprefix("TX:") for value in (existing_ids or set())}
    existing_events = existing_events or set()
    records, held = [], []
    for row in rows:
        fields = row["fields"]
        employer = str(fields["JOB_SITE_NAME"] or "").strip()
        city = str(fields["CITY_NAME"] or "").strip()
        county = str(fields["COUNTY_NAME"] or "").strip()
        notice = row["notice_date"]
        action = row["effective_date"]
        identity = f"TX:{row['source_notice_id']}"
        event = (employer.casefold(), notice, city.casefold(), county.casefold())
        reason = ("anomalous_action_year" if action and int(action[:4]) > int(notice[:4]) + 2 else
                  "missing_employer" if not employer else
                  "missing_location" if not city or not county else
                  "already_represented_tx_id" if row["source_notice_id"] in existing_ids else
                  "idless_current_event_overlap" if event in existing_events else None)
        if reason:
            held.append({"origin": SOURCE_NAME, "state": "TX", "reason": reason,
                         "source_row": row["source_row"], "source_row_sha256": row["row_sha256"],
                         "source_notice_id": row["source_notice_id"],
                         "source_url": manifest["source_url"], "notice_year": notice[:4],
                         "raw_extra": _json(fields)})
            continue
        details = {"origin": SOURCE_NAME, "source_row": row["source_row"],
                   "source_row_sha256": row["row_sha256"],
                   "source_artifact_sha256": row["artifact_sha256"],
                   "identity_basis": "TWC_NOTICE_ID",
                   "agency_received_date": row["agency_received_date"],
                   "date_roles": {"NOTICE_DATE": "reported_notice",
                                  "LayOff_Date": "reported_action",
                                  "WFDD_RECEIVED_DATE": "agency_receipt"},
                   "raw_fields": fields}
        rec = {"state": "TX", "employer_name": employer,
               "location": f"{city}, {county}", "notice_date": notice,
               "notice_date_precision": "day", "notice_date_basis": "reported",
               "effective_date": action,
               "effective_date_precision": "day" if action else None,
               "effective_date_basis": "reported" if action else None,
               "employees_affected": row["workers"], "layoff_type": "unknown",
               "is_temporary": fields["Temporary_Layoff_Flag"], "is_amendment": 0,
               "source_url": manifest["source_url"],
               "source_notice_id": row["source_notice_id"],
               "source_identity": identity, "source_details": _json(details),
               "raw_extra": _json(fields),
               "dedupe_key": hashlib.sha1(f"TX|source|{identity}".encode()).hexdigest()}
        rec["raw_record_hash"] = _record_hash(rec)
        records.append(rec)
    if len(records) + len(held) != EXPECTED_ROWS:
        raise ValueError("Texas historical source row accounting mismatch")
    return records, held, {"source_rows": EXPECTED_ROWS, "admitted": len(records),
                           "held": len(held),
                           "hold_reasons": dict(Counter(item["reason"] for item in held)),
                           "admitted_missing_workers": sum(rec["employees_affected"] is None for rec in records),
                           "admitted_missing_action_date": sum(rec["effective_date"] is None for rec in records)}


def build(base_bundle: Path, artifacts: Path, out_bundle: Path) -> dict:
    """Overlay the exact pinned workbook on an agency-only replay bundle."""
    base_bundle, artifacts, out_bundle = map(Path, (base_bundle, artifacts, out_bundle))
    if out_bundle.exists():
        raise FileExistsError(out_bundle)
    base = verify(base_bundle)
    if base.get("admission_inputs") != "agency-only-v1":
        raise ValueError("Texas overlay requires an agency-only base")
    read_artifacts(artifacts)
    with tarfile.open(base_bundle, "r:gz") as archive:
        files = {member.name: archive.extractfile(member).read()
                 for member in archive.getmembers() if member.name != "manifest.json"}
    if any(name.startswith("agency/tx_historical/") for name in files):
        raise ValueError("base bundle already contains Texas historical workbook")
    for name in ("manifest.json", "tx_historical.xlsx"):
        files[f"agency/tx_historical/{name}"] = (artifacts / name).read_bytes()
    manifest = {**base, "files": [
        {"path": name, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        for name, content in sorted(files.items())]}
    out_bundle.parent.mkdir(parents=True, exist_ok=True)
    try:
        with out_bundle.open("xb") as raw, gzip.GzipFile(fileobj=raw, mode="wb", filename="", mtime=0) as zipped, tarfile.open(fileobj=zipped, mode="w") as archive:
            header = (json.dumps(manifest, sort_keys=True, indent=2) + "\n").encode()
            archive.addfile(_entry("manifest.json", header), io.BytesIO(header))
            for name, content in sorted(files.items()):
                archive.addfile(_entry(name, content), io.BytesIO(content))
        verify(out_bundle)
    except BaseException:
        out_bundle.unlink(missing_ok=True)
        raise
    return manifest
