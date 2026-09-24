"""Conservative projection of Missouri's agency-published rapid-response workbook.

The fiscal-year sheets are a mixed service register, not a list of 1,301
verified WARN filings. Admission requires an individually reviewed statement
that an actual WARN was issued, received, or present. Each approval is pinned
to the worksheet, physical row, and exact raw-cell hash.
"""

from __future__ import annotations

import gzip
import hashlib
import io
import json
import tarfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from warnlive.migrate.source_bundle import _entry, verify
from warnlive.normalize.engine import _record_hash

ARTIFACT = "WARN_Data1997-2018.xlsx"
PREFIX = "agency/mo_historical"
EXPECTED_SHA256 = "3e783842d6086e7f176c20d81de6906873c1bb2903126a3ceee900adece8141c"

# Reviewed comments say a WARN was issued/received, or refer to the actual
# WARN's contents. The raw-cell hash makes row-number reuse fail closed.
REVIEWED_WARN_ROWS = {
    ("2005-06 Data", 39, "0b89c86510dea6cf35bb1cba41661145fc91facb57da808eaf5200f12d6649c2"),
    ("2005-06 Data", 47, "0749f4a1142be9d3e6956f58bb9de5490a7f2418dcbfab52af0c9f15283dc875"),
    ("2005-06 Data", 49, "c746906f5b53308927ffee4331b3b4008dcb24b1925226eeee41b16bb16480ec"),
    ("2007-08 Data", 24, "2a13e63223298d25476f40ed29d9bdd11457dba28b59c5b5d037bb3132e1e635"),
    ("2008-09 Data", 34, "ee717a6e60e691275e01ae53f5a1677a248e2a50dec41d9d5c1799fa9e843a29"),
    ("2008-09 Data", 35, "3a4a0b6c70cf5d9ee578e3fbae7e2197f9d3e1a3406dd7c92c92d25cb7beeb21"),
    ("2008-09 Data", 36, "550a9e5b1d61c0d7b77fb5b0cf29e70d7657aecae09a57da47d451fa281e8454"),
    ("2008-09 Data", 37, "51e4ddfb81b425e48304ca34f7b7726eb672ea82bad057a4dd70da55e5773ca3"),
    ("2008-09 Data", 44, "01719482053a243844e4d891111bf774638e63b399c151e3cd657a00bdac0519"),
    ("2008-09 Data", 54, "67551a8cb3af80fa6637e4c1f0ec357026d2bb81487798a7cdd08281183d70fc"),
    ("2010-11 Data", 13, "7ad104c2b0d7803665c81df0ccf79e56cd8747fa83ff6335377939dd4b4a0471"),
    ("2010-11 Data", 21, "c143f44048083410fa9667acdc16ea90d03caa2e28799e68538f094e217ad316"),
    ("2011-12 Data", 8, "ec6124590cb44f0ed51316c9bdc29cbdb19e418fba602836a96b423dfc633631"),
    ("2011-12 Data", 19, "d135f49b45eab37233bad98db31ce9131e368d90fa58ed006a1e49180da9ee80"),
    ("2011-12 Data", 22, "f26c6dfc2c5431dd0a41376534361b75845d4d85a21f5fccdc873b9fd07e7fcb"),
    ("2012-13 Data", 20, "88863ad33b8451162d620b2b15bc350d47888e7f362d1e962bac4d2a9b502671"),
    ("2017-18 Data", 19, "90234b1470f78eb45b199dce7fd5d5962bc98ac484770ecc866b8ce9e23e385f"),
}
UNCERTAIN_ACTION_ROWS = {("2012-13 Data", 20)}  # Penske comment says date moved forward.


def _json(value: object) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, default=str,
                      separators=(",", ":"))


def _raw(value: object) -> object:
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _day(value: object) -> str | None:
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return None
    if isinstance(value, datetime) and value.time() == datetime.min.time():
        return value.date().isoformat()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    return None


def read_artifacts(directory: Path) -> tuple[list[dict], dict]:
    """Verify source bytes and manifest accounting; retain every company row."""
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    path = directory / ARTIFACT
    if (manifest.get("format") != "agency-linked-research-workbook-v1"
            or manifest.get("file") != ARTIFACT
            or manifest.get("sha256") != EXPECTED_SHA256
            or manifest.get("sheet_count") != 22
            or manifest.get("data_rows_with_company") != 1301
            or path.is_symlink()):
        raise ValueError("unsupported Missouri historical manifest")
    content = path.read_bytes()
    if (len(content) != manifest.get("bytes")
            or hashlib.sha256(content).hexdigest() != EXPECTED_SHA256):
        raise ValueError("Missouri historical workbook checksum mismatch")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if book.sheetnames != [item["name"] for item in manifest["sheets"]]:
            raise ValueError("Missouri historical sheets changed")
        rows = []
        for sheet, spec in zip(book.worksheets, manifest["sheets"], strict=True):
            iterator = sheet.values
            headers = tuple(next(iterator))
            if (len(headers) != len(set(headers)) or "Company Name" not in headers
                    or "Layoff or Closing Date" not in headers
                    or "# Affected" not in headers or "Comments" not in headers):
                raise ValueError(f"Missouri historical headers changed: {sheet.title}")
            company_rows = []
            for physical_row, cells in enumerate(iterator, start=2):
                if len(cells) != len(headers):
                    raise ValueError(f"Missouri historical row width changed: {sheet.title}:{physical_row}")
                raw = {key: _raw(value) for key, value in zip(headers, cells, strict=True)}
                if not str(raw["Company Name"] or "").strip():
                    continue
                item = {"sheet": sheet.title, "physical_row": physical_row,
                        "raw": raw,
                        "source_row_sha256": hashlib.sha256(_json(raw).encode()).hexdigest()}
                company_rows.append(item)
                rows.append(item)
            if (len(company_rows) != spec["data_rows_with_company"]
                    or company_rows[0]["physical_row"] != spec["first_data_row"]
                    or company_rows[-1]["physical_row"] != spec["last_data_row"]):
                raise ValueError(f"Missouri historical row accounting changed: {sheet.title}")
        if len(rows) != 1301:
            raise ValueError("Missouri historical total row accounting changed")
        return rows, manifest
    finally:
        book.close()


def project(directory: Path, existing_events: set[tuple[str, str | None]] | None = None
            ) -> tuple[list[dict], list[dict], dict]:
    """Admit only reviewed actual-WARN rows, retaining all other rows as holds.

    ``existing_events`` is a caller-supplied set of (casefolded employer,
    reported action date) pairs from current agency sources. It is a cautious
    overlap screen, not an assertion that the pair is a filing identity.
    """
    rows, manifest = read_artifacts(directory)
    approved = set(REVIEWED_WARN_ROWS)
    found = {(r["sheet"], r["physical_row"], r["source_row_sha256"]) for r in rows}
    if not approved <= found:
        raise ValueError("reviewed Missouri WARN row is absent or changed")
    existing_events = existing_events or set()
    records, held = [], []
    for row in rows:
        raw = row["raw"]
        sheet, physical_row, row_sha = (row["sheet"], row["physical_row"],
                                        row["source_row_sha256"])
        reviewed = (sheet, physical_row, row_sha) in approved
        employer = str(raw["Company Name"] or "").strip()
        location = str(raw["Location(s)"] or "").strip()
        county = str(raw["County"] or "").strip()
        received = _day(next(iter(raw.values())))
        action = _day(raw["Layoff or Closing Date"])
        action_conflict = (sheet, physical_row) in UNCERTAIN_ACTION_ROWS
        workers = raw["# Affected"]
        valid_workers = (isinstance(workers, int) and not isinstance(workers, bool)
                         and workers >= 0)
        location_text = ", ".join(part for part in (location, county) if part)
        pointer = (f"{PREFIX}/{ARTIFACT}:sha256:{manifest['sha256']}:"
                   f"sheet:{sheet}:row:{physical_row}")
        reason = ("unreviewed_rapid_response_row" if not reviewed else
                  "missing_agency_receipt_date" if not received else
                  "compound_or_missing_action_date" if not action and not action_conflict else
                  "compound_or_missing_worker_count" if not valid_workers else
                  "missing_location" if not location_text else
                  "current_agency_event_overlap" if (employer.casefold(), action) in existing_events else None)
        if reason:
            held.append({"origin": f"{PREFIX}/{ARTIFACT}", "state": "MO",
                         "reason": reason, "source_row": pointer,
                         "source_row_sha256": row_sha, "source_sheet": sheet,
                         "physical_row": physical_row, "source_notice_id": None,
                         "source_url": manifest["source_url"],
                         "notice_year": received[:4] if received else None,
                         "raw_extra": _json(raw)})
            continue
        details = {"origin": f"{PREFIX}/{ARTIFACT}",
                   "source_row": pointer, "source_row_sha256": row_sha,
                   "source_artifact_sha256": manifest["sha256"],
                   "identity_basis": "reviewed_source_row_observation_no_filing_id",
                   "admission_basis": "comment_attests_actual_WARN",
                   "agency_received_date": received,
                   "action_date_conflict": ("original WARN date moved forward; actual separation date absent"
                                            if action_conflict else None),
                   "date_roles": {"Date Rec'd": "agency_receipt",
                                  "Layoff or Closing Date": "reported_action"},
                   "raw_fields": raw}
        kind = str(raw["Type of Notice"] or "").lower()
        record = {"state": "MO", "employer_name": employer,
                  "location": location_text, "notice_date": None,
                  "effective_date": None if action_conflict else action,
                  "effective_date_precision": None if action_conflict else "day",
                  "effective_date_basis": None if action_conflict else "reported",
                  "employees_affected": workers,
                  "layoff_type": "closure" if "clos" in kind else
                                 "mass_layoff" if "layoff" in kind else "unknown",
                  "is_temporary": None, "is_amendment": 0,
                  "source_url": manifest["source_url"],
                  "source_notice_id": None, "source_identity": None,
                  "source_details": _json(details), "raw_extra": _json(raw),
                  # Stable source observation anchor; not a fabricated filing ID.
                  "dedupe_key": hashlib.sha1(
                      f"MO|agency-workbook|{sheet}|{physical_row}|{row_sha}".encode()
                  ).hexdigest()}
        record["raw_record_hash"] = _record_hash(record)
        records.append(record)
    if len(records) + len(held) != 1301:
        raise ValueError("Missouri historical source row accounting mismatch")
    return records, held, {"source_rows": 1301, "admitted": len(records),
                           "held": len(held),
                           "hold_reasons": dict(Counter(item["reason"] for item in held))}


def build(base_bundle: Path, artifacts: Path, out_bundle: Path) -> dict:
    """Overlay the pinned agency workbook on an agency-only replay bundle."""
    base_bundle, artifacts, out_bundle = map(Path, (base_bundle, artifacts, out_bundle))
    if out_bundle.exists():
        raise FileExistsError(out_bundle)
    base = verify(base_bundle)
    if base.get("admission_inputs") != "agency-only-v1":
        raise ValueError("Missouri historical overlay requires an agency-only base")
    read_artifacts(artifacts)
    with tarfile.open(base_bundle, "r:gz") as archive:
        files = {member.name: archive.extractfile(member).read()
                 for member in archive.getmembers() if member.name != "manifest.json"}
    if any(name.startswith(PREFIX + "/") for name in files):
        raise ValueError("base bundle already contains Missouri historical source")
    for name in ("manifest.json", ARTIFACT):
        files[f"{PREFIX}/{name}"] = (artifacts / name).read_bytes()
    manifest = {**base, "files": [
        {"path": name, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        for name, content in sorted(files.items())]}
    out_bundle.parent.mkdir(parents=True, exist_ok=True)
    try:
        with out_bundle.open("xb") as raw, gzip.GzipFile(fileobj=raw, mode="wb", filename="", mtime=0) as zipped, tarfile.open(fileobj=zipped, mode="w") as archive:
            header = (json.dumps(manifest, sort_keys=True, indent=2) + "\n").encode()
            archive.addfile(_entry("manifest.json", header), io.BytesIO(header))
            for name, content in sorted(files.items()):
                archive.addfile(_entry(name, content), io.BytesIO(content))
        verify(out_bundle)
    except BaseException:
        out_bundle.unlink(missing_ok=True)
        raise
    return manifest
