"""Project the pinned Kentucky Career Center reports without inventing notice dates.

``read_artifacts``/``project``/``build`` handle ``agency/ky``, the linked 2026
report CSV. ``read_archive``/``project_archive`` handle ``agency/ky_archive``
(data/source_snapshots/ky/kcc-2026-09-30): the 1998-2016 tracking workbook,
the 2017-Feb 2025 report workbook and the 2025 report CSV. Each linked file is
a snapshot, not a complete annual inventory. Notice numbers, else notice
document URLs, else the workbook row identify agency rows; ``Date Received``
is an agency receipt day, while ``Projected Date(s)`` describes the planned
action. No network access.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import re
import tarfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from warnlive.migrate.source_bundle import _entry, verify
from warnlive.normalize.engine import _record_hash


PREFIX = "agency/ky"
FILENAME = "WARN-Report-2026-09-16.csv"
SHA256 = "a40a050d1b10a2ea4d1ec54302d548c4499534d1e16411840b9c383d5d2b2da5"
HEADERS = (
    "Company: Company Name", "Notice Type", "Notice: Notice Number",
    "Closure or Layoff?", "County", "Date Received", "NAICS",
    "Notice URL", "Number of Employees Affected", "Projected Date",
    "Trade", "Type of Employees Affected", "Workforce Board",
)
NOTICE_ID = re.compile(r"Notice ([1-9][0-9]*)\Z")
COUNT = 35
IN_STATE = 33
OUT_OF_STATE = {"2833", "2778"}


def _json(value: object) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _date(value: str, field: str, row_number: int) -> str:
    if not re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", value):
        raise ValueError(f"invalid Kentucky {field} at row {row_number}: {value!r}")
    try:
        return datetime.strptime(value, "%m/%d/%Y").date().isoformat()
    except ValueError as exc:
        raise ValueError(f"invalid Kentucky {field} at row {row_number}: {value!r}") from exc


def _artifact(directory: Path, name: str) -> Path:
    path = directory / name
    if directory.is_symlink() or path.is_symlink() or not path.is_file():
        raise ValueError(f"invalid Kentucky source artifact: {path}")
    return path


def read_artifacts(directory: Path) -> list[dict]:
    """Verify source bytes and every physical CSV row against the pinned manifest."""
    directory = Path(directory)
    manifest = json.loads(_artifact(directory, "manifest.json").read_text())
    if (manifest.get("format") != "ky-agency-report-capture-v1"
            or manifest.get("file") != FILENAME
            or manifest.get("sha256") != SHA256
            or manifest.get("rows") != COUNT
            or manifest.get("in_state_rows") != IN_STATE
            or manifest.get("unique_notice_numbers") != COUNT
            or manifest.get("observed_utc_date") != "2026-09-24"
            or manifest.get("source_page") != "https://kcc.ky.gov/Pages/News.aspx"
            or manifest.get("source_url") != (
                "https://kcc.ky.gov/WARN%20notices/WARN%20Notices%202026/"
                "WARN%20Report-2026-09-16-09-00-01.csv")
            or not isinstance(manifest.get("out_of_state_hold_rows"), list)):
        raise ValueError("Kentucky source manifest changed")
    path = _artifact(directory, FILENAME)
    content = path.read_bytes()
    if len(content) != manifest.get("bytes") or hashlib.sha256(content).hexdigest() != SHA256:
        raise ValueError("Kentucky source checksum mismatch")
    with io.StringIO(content.decode("utf-8-sig"), newline="") as stream:
        reader = csv.reader(stream)
        if tuple(next(reader, ())) != HEADERS:
            raise ValueError("Kentucky source header changed")
        observations = []
        ids: set[str] = set()
        urls: set[str] = set()
        for row_number, cells in enumerate(reader, start=2):
            if len(cells) != len(HEADERS):
                raise ValueError(f"Kentucky CSV row width changed: {row_number}")
            raw = dict(zip(HEADERS, cells, strict=True))
            match = NOTICE_ID.fullmatch(raw["Notice: Notice Number"])
            if not match or match.group(1) in ids:
                raise ValueError(f"invalid or duplicate Kentucky notice ID at row {row_number}")
            ident = match.group(1)
            ids.add(ident)
            document = raw["Notice URL"]
            parsed = urlparse(document)
            if (parsed.scheme != "https" or parsed.netloc != "kydev.my.salesforce.com"
                    or not parsed.path.startswith("/sfc/p/") or document in urls):
                raise ValueError(f"invalid or duplicate Kentucky notice URL at row {row_number}")
            urls.add(document)
            if (not raw["Company: Company Name"].strip()
                    or raw["Notice Type"] != "WARN Notice"
                    or not raw["County"].strip()
                    or raw["Closure or Layoff?"] not in {"", "Closure", "Layoff"}
                    or not raw["Number of Employees Affected"].isdigit()
                    or int(raw["Number of Employees Affected"]) <= 0):
                raise ValueError(f"incomplete Kentucky source row {row_number}")
            received = _date(raw["Date Received"], "Date Received", row_number)
            projected = _date(raw["Projected Date"], "Projected Date", row_number)
            if received[:4] not in {"2025", "2026"} or projected[:4] != "2026":
                raise ValueError(f"Kentucky source date scope changed at row {row_number}")
            raw_hash = hashlib.sha256(_json(raw).encode()).hexdigest()
            observations.append({
                "source_artifact": f"{PREFIX}/{FILENAME}",
                "source_artifact_sha256": SHA256,
                "source_row": f"{PREFIX}/{FILENAME}:sha256:{SHA256}:row:{row_number}:sha256:{raw_hash}",
                "source_row_sha256": raw_hash,
                "source_url": manifest["source_url"],
                "document_url": document,
                "row_number": row_number,
                "source_notice_id": ident,
                "agency_received_date": received,
                "projected_date": projected,
                "workers_reported": int(raw["Number of Employees Affected"]),
                "company_text": raw["Company: Company Name"],
                "county_text": raw["County"],
                "effective_date": projected,
                "raw": raw,
                "raw_fields": raw,
            })
    if len(observations) != COUNT:
        raise ValueError("Kentucky source row count changed")
    held = [row for row in observations if row["raw_fields"]["County"] == "Out of the State County"]
    expected_holds = {(
        item.get("notice_id"), item.get("employer"), item.get("county"), item.get("employees"))
        for item in manifest["out_of_state_hold_rows"]}
    actual_holds = {(
        f"Notice {row['source_notice_id']}", row["raw_fields"]["Company: Company Name"],
        row["raw_fields"]["County"], row["raw_fields"]["Number of Employees Affected"])
        for row in held}
    if (len(held) != COUNT - IN_STATE or
            {row["source_notice_id"] for row in held} != OUT_OF_STATE or
            expected_holds != actual_holds):
        raise ValueError("Kentucky out-of-state row boundary changed")
    return observations


def project(directory: Path) -> tuple[list[dict], list[dict], dict]:
    """Admit in-state agency IDs and account for the two out-of-state rows."""
    rows = read_artifacts(directory)
    admitted: list[dict] = []
    held: list[dict] = []
    for row in rows:
        raw = row["raw_fields"]
        if raw["County"] == "Out of the State County":
            held.append({
                "origin": row["source_artifact"], "state": "KY",
                "reason": "source_explicitly_out_of_state",
                "source_row": row["source_row"],
                "source_row_sha256": row["source_row_sha256"],
                "source_notice_id": row["source_notice_id"],
                "source_url": row["source_url"],
                "document_url": row["document_url"],
                "notice_year": row["agency_received_date"][:4],
                "raw_extra": _json(raw),
            })
            continue
        identity = f"KY:{row['source_notice_id']}"
        kind = raw["Closure or Layoff?"]
        details = {
            "source_artifact": row["source_artifact"],
            "source_artifact_sha256": row["source_artifact_sha256"],
            "source_row": row["source_row"],
            "source_row_sha256": row["source_row_sha256"],
            "source_document_url": row["document_url"],
            "identity_basis": "agency_notice_number",
            "agency_received_date": row["agency_received_date"],
            "projected_action_date": row["projected_date"],
            "date_roles": {"Date Received": "agency_receipt",
                           "Projected Date": "projected_action"},
            "date_evidence_rule": "ky_agency_report_labeled_fields_v1",
            "location_basis": "agency_county_field",
            "raw_fields": raw,
        }
        record = {
            "state": "KY", "employer_name": raw["Company: Company Name"].strip(),
            "location": raw["County"].strip(),
            "notice_date": None,
            "effective_date": row["projected_date"],
            "effective_date_precision": "day", "effective_date_basis": "reported",
            "employees_affected": row["workers_reported"],
            "layoff_type": "closure" if kind == "Closure" else "mass_layoff" if kind == "Layoff" else "unknown",
            "is_temporary": None, "is_amendment": 0,
            "source_url": row["source_url"],
            "source_notice_id": row["source_notice_id"],
            "source_identity": identity,
            "source_details": _json(details),
            "raw_extra": _json(raw),
            "dedupe_key": hashlib.sha1(f"KY|source|{identity}".encode()).hexdigest(),
        }
        record["raw_record_hash"] = _record_hash(record)
        admitted.append(record)
    if len(admitted) != IN_STATE or len(admitted) + len(held) != COUNT:
        raise ValueError("Kentucky source row accounting mismatch")
    return admitted, held, {
        "source_rows": COUNT, "admitted": len(admitted), "held": len(held),
        "hold_reasons": dict(Counter(item["reason"] for item in held)),
        "admitted_workers": sum(item["employees_affected"] for item in admitted),
    }


def build(base_bundle: Path, artifacts: Path, out_bundle: Path) -> dict:
    """Add the exact pinned agency CSV and manifest to an agency-only bundle."""
    base_bundle, artifacts, out_bundle = map(Path, (base_bundle, artifacts, out_bundle))
    if out_bundle.exists():
        raise FileExistsError(out_bundle)
    base = verify(base_bundle)
    if base.get("admission_inputs") != "agency-only-v1":
        raise ValueError("Kentucky overlay requires an agency-only base")
    read_artifacts(artifacts)
    with tarfile.open(base_bundle, "r:gz") as archive:
        files = {member.name: archive.extractfile(member).read()
                 for member in archive.getmembers() if member.name != "manifest.json"}
    if any(name.startswith(PREFIX + "/") for name in files):
        raise ValueError("base bundle already contains Kentucky source")
    if "raw/ky.csv" in files:
        raise ValueError("base bundle still contains unverified Kentucky raw capture")
    for name in ("manifest.json", FILENAME):
        files[f"{PREFIX}/{name}"] = (artifacts / name).read_bytes()
    manifest = {**base, "files": [
        {"path": name, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        for name, content in sorted(files.items())]}
    out_bundle.parent.mkdir(parents=True, exist_ok=True)
    try:
        with out_bundle.open("xb") as raw, gzip.GzipFile(
            fileobj=raw, mode="wb", filename="", mtime=0
        ) as zipped, tarfile.open(fileobj=zipped, mode="w") as archive:
            header = (json.dumps(manifest, sort_keys=True, indent=2) + "\n").encode()
            archive.addfile(_entry("manifest.json", header), io.BytesIO(header))
            for name, content in sorted(files.items()):
                archive.addfile(_entry(name, content), io.BytesIO(content))
        verify(out_bundle)
    except BaseException:
        out_bundle.unlink(missing_ok=True)
        raise
    return manifest


# ---------------------------------------------------------------------------
# Kentucky agency archive (agency/ky_archive): the 1998-2016 tracking workbook,
# the 2017-Feb 2025 report workbook and the 2025 report CSV, captured live from
# kcc.ky.gov on 2026-09-30 (data/source_snapshots/ky/kcc-2026-09-30).

ARCHIVE_PREFIX = "agency/ky_archive"
ARCHIVE_FORMAT = "ky-agency-archive-v1"
TRACKING = "tracking-form-1998-2016.xlsx"
REPORT = "warn-report-2017-2025-02.xlsx"
REPORT_2025 = "warn-report-2025.csv"
ARCHIVE_SHA256 = {
    TRACKING: "2411db199ba2eefc3b39659b74fc5c4b1775f9c013ad6478dcb6e7a665285c78",
    REPORT: "8067e0bdf3358a081cdccd15ad5a18b04b09063720787a6adbfa858943c120f0",
    REPORT_2025: "d7ae13baebd26f220a9b0d38dc08f3d26919225a0b63621cb05baaba32ba282b",
}
ARCHIVE_ROWS = {TRACKING: 799, REPORT: 368, REPORT_2025: 60}
REPORT_HEADERS = (
    "Date Received", "Region", "County", "Company Name", "NAICS Code", "Employees",
    "Closure or Layoff?", "Projected Date", "Trade", "Notice URL", "Notice Link",
)
# A comment or employer cell that declares the row an amendment, revision or
# rescission of another notice. The tracking form never names the row it
# amends, so these rows are held rather than linked.
AMENDMENT = re.compile(r"\bamend|\brevised?\b|\bre-notice\b|continuation of previous", re.I)
RESCINDED = re.compile(r"\brescind", re.I)
SALESFORCE_DOC = re.compile(r"https://kydev\.my\.salesforce\.com/sfc/p/[^/]+/a/([A-Za-z0-9]+)/\S+")
OUT_OF_STATE_TEXT = re.compile(r"\bout of (?:the )?state\b", re.I)


def _cell(value: object) -> object:
    """A JSON-safe copy of one workbook cell value."""
    if isinstance(value, datetime):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _text(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def _serial_date(value: object) -> str | None:
    """A whole-cell date: a date-typed cell (kept as its ISO text), an Excel
    serial or m/d/yyyy text."""
    from openpyxl.utils.datetime import from_excel

    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}", value):
        return value[:10]
    if isinstance(value, int) and not isinstance(value, bool) and 30000 <= value <= 60000:
        return from_excel(value).date().isoformat()
    text = _text(value)
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", text):
        try:
            return datetime.strptime(text, "%m/%d/%Y").date().isoformat()
        except ValueError:
            return None
    return None


def _date_span(value: object) -> tuple[str | None, str | None]:
    """A single projected date, or both ends of an explicit ``a - b`` range."""
    single = _serial_date(value)
    if single:
        return single, None
    match = re.fullmatch(r"(\d{1,2}/\d{1,2}/\d{4})\s*[-–]\s*(\d{1,2}/\d{1,2}/\d{4})", _text(value))
    if not match:
        return None, None
    start, end = (_serial_date(part) for part in match.groups())
    if not start or not end or end < start:
        return None, None
    return start, end


def _workers(value: object, stored_serial: object = None) -> int | None:
    """A whole-cell positive count; anything else (ranges, "+/-", notes) is blank."""
    if stored_serial is not None:
        value = stored_serial
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        return value if value > 0 else None
    text = _text(value).replace(",", "")
    return int(text) if text.isdigit() and int(text) > 0 else None


def _archive_artifact(directory: Path, name: str, spec: dict) -> bytes:
    path = _artifact(directory, name)
    content = path.read_bytes()
    if (spec.get("sha256") != ARCHIVE_SHA256[name] or len(content) != spec.get("bytes")
            or hashlib.sha256(content).hexdigest() != ARCHIVE_SHA256[name]):
        raise ValueError(f"Kentucky archive checksum mismatch: {name}")
    return content


def _workbook_rows(content: bytes, name: str) -> list[dict]:
    """Every non-blank data row of every sheet, keyed by its cleaned header."""
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel

    book = load_workbook(io.BytesIO(content), data_only=True)
    rows = []
    for sheet in book.worksheets:
        cells = list(sheet.iter_rows())
        header = [_text(cell.value) or f"column_{index}" for index, cell in
                  enumerate(cells[0], start=1)]
        if "Company Name" not in header or len(set(header)) != len(header):
            raise ValueError(f"Kentucky archive header changed: {name}:{sheet.title}")
        for row in cells[1:]:
            if not any(_text(cell.value) for cell in row):
                continue
            raw = {head: _cell(cell.value) for head, cell in zip(header, row)
                   if cell.value is not None and _text(cell.value) != ""}
            # A count typed into a date-formatted cell is stored as that
            # number; the workbook only displays it as a 1900 date.
            stored = {}
            for head, cell in zip(header, row):
                if head == "Employees" and hasattr(cell.value, "year"):
                    serial = to_excel(cell.value)
                    stored[head] = int(serial) if float(serial).is_integer() else None
                elif head == "Employees" and hasattr(cell.value, "hour"):
                    stored[head] = 0
            rows.append({"sheet": sheet.title, "row_number": row[0].row,
                         "header": header, "raw": raw, "stored_numbers": stored})
    return rows


def read_archive(directory: Path) -> list[dict]:
    """Verify the pinned archive and return one observation per source row."""
    directory = Path(directory)
    manifest = json.loads(_artifact(directory, "manifest.json").read_text())
    specs = {item.get("file"): item for item in manifest.get("artifacts") or []}
    if (manifest.get("format") != ARCHIVE_FORMAT or manifest.get("state") != "KY"
            or set(specs) != set(ARCHIVE_SHA256)
            or any(specs[name].get("data_rows") != count for name, count in ARCHIVE_ROWS.items())):
        raise ValueError("Kentucky archive manifest changed")
    observations = []
    for name in (TRACKING, REPORT):
        spec = specs[name]
        rows = _workbook_rows(_archive_artifact(directory, name, spec), name)
        counts = Counter(row["sheet"] for row in rows)
        if len(rows) != ARCHIVE_ROWS[name] or dict(counts) != spec.get("sheet_data_rows"):
            raise ValueError(f"Kentucky archive row count changed: {name}")
        if name == REPORT and any(tuple(row["header"][:11]) != REPORT_HEADERS for row in rows):
            raise ValueError("Kentucky report workbook header changed")
        for row in rows:
            payload = {"sheet": row["sheet"], "row": row["row_number"], "cells": row["raw"],
                       "stored_numbers": row["stored_numbers"]}
            raw_hash = hashlib.sha256(_json(payload).encode()).hexdigest()
            observations.append({
                "file": name, "sheet": row["sheet"], "row_number": row["row_number"],
                "raw": row["raw"], "stored_numbers": row["stored_numbers"],
                "source_artifact": f"{ARCHIVE_PREFIX}/{name}",
                "source_artifact_sha256": ARCHIVE_SHA256[name],
                "source_row": (f"{ARCHIVE_PREFIX}/{name}:sha256:{ARCHIVE_SHA256[name]}:"
                               f"sheet:{row['sheet']}:row:{row['row_number']}"),
                "source_row_sha256": raw_hash,
                "source_url": spec["original_agency_url"],
            })
    spec = specs[REPORT_2025]
    content = _archive_artifact(directory, REPORT_2025, spec)
    with io.StringIO(content.decode("utf-8-sig"), newline="") as stream:
        reader = csv.reader(stream)
        if tuple(next(reader, ())) != HEADERS:
            raise ValueError("Kentucky 2025 report header changed")
        count = 0
        for row_number, cells in enumerate(reader, start=2):
            if len(cells) != len(HEADERS):
                raise ValueError(f"Kentucky 2025 report row width changed: {row_number}")
            raw = dict(zip(HEADERS, cells, strict=True))
            raw_hash = hashlib.sha256(_json(raw).encode()).hexdigest()
            count += 1
            observations.append({
                "file": REPORT_2025, "sheet": None, "row_number": row_number, "raw": raw,
                "stored_numbers": {},
                "source_artifact": f"{ARCHIVE_PREFIX}/{REPORT_2025}",
                "source_artifact_sha256": ARCHIVE_SHA256[REPORT_2025],
                "source_row": (f"{ARCHIVE_PREFIX}/{REPORT_2025}:sha256:"
                               f"{ARCHIVE_SHA256[REPORT_2025]}:row:{row_number}:sha256:{raw_hash}"),
                "source_row_sha256": raw_hash,
                "source_url": spec["original_agency_url"],
            })
    if count != ARCHIVE_ROWS[REPORT_2025]:
        raise ValueError("Kentucky 2025 report row count changed")
    return observations


def _archive_fields(row: dict) -> dict:
    """Map one archive row onto the shared field names, keeping source text."""
    raw, name = row["raw"], row["file"]
    if name == REPORT_2025:
        number = NOTICE_ID.fullmatch(raw["Notice: Notice Number"])
        received = _serial_date(raw["Date Received"])
        start, end = _date_span(raw["Projected Date"])
        return {"company": _text(raw["Company: Company Name"]), "county": _text(raw["County"]),
                "address": None, "workers": _workers(raw["Number of Employees Affected"]),
                "kind": raw["Closure or Layoff?"], "received": received,
                "projected": (start, end), "projected_field": "Projected Date",
                "document_url": raw["Notice URL"] or None,
                "notice_id": number.group(1) if number else None,
                "year": (received or "")[:4] or None, "comment": "", "source": "WARN"}
    received = _serial_date(raw.get("Date Received"))
    projected_field = "Projected Date" if name == REPORT else "Projected Dates"
    county = raw.get("County") if name == REPORT else raw.get("County/Counties", raw.get("County"))
    address = None if name == REPORT else raw.get("Location", raw.get("Company Address"))
    sheet_year = re.search(r"(\d{4})", row["sheet"]).group(1)
    return {"company": _text(raw.get("Company Name")), "county": _text(county) or None,
            "address": _text(address) or None,
            "workers": _workers(raw.get("Employees"), row["stored_numbers"].get("Employees")),
            "kind": _text(raw.get("Closure or Layoff?", raw.get("Closure/Layoff"))),
            "received": received, "projected": _date_span(raw.get(projected_field)),
            "projected_field": projected_field,
            "document_url": _text(raw.get("Notice URL")) or None, "notice_id": None,
            "year": sheet_year, "comment": _text(raw.get("Comments")),
            "source": _text(raw.get("Source")) if "Source" in raw else None}


def _occupations_match(a: object, b: object) -> bool:
    """Equal occupation text, or both only a pointer to the notice ("See WARN").

    A named occupation ("Senior Engineer") distinguishes one-worker listings
    of one employer and day, so such rows are kept apart.
    """
    a, b = _text(a).casefold(), _text(b).casefold()
    return a == b or (a.startswith("see ") and b.startswith("see "))


def _reentry_key(row: dict, f: dict) -> str | None:
    """The facts a re-typed tracking-sheet listing repeats; None off the sheet.

    Employer, county, address, Date Received, projected action date (parsed
    start, else the cell text) and employee count, normalized.
    """
    if row["file"] != TRACKING or not f["company"]:
        return None
    fold = lambda value: re.sub(r"[^a-z0-9]", "", str(value or "").casefold())  # noqa: E731
    projected = f["projected"][0] or _text(row["raw"].get(f["projected_field"])).casefold()
    return _json([fold(f["company"]), fold(f["county"]), fold(f["address"]), f["received"],
                  projected, f["workers"]])


def project_archive(directory: Path, existing_ids: set[str] | None = None,
                    existing_urls: set[str] | None = None) -> tuple[list[dict], list[dict], dict]:
    """Admit archive rows as notices; hold duplicates, amendments and non-WARN rows.

    ``existing_ids`` and ``existing_urls`` are the notice numbers and notice
    document URLs of the Kentucky rows already in the build (agency/ky). An
    agency notice number is the identity where the source has one; otherwise
    the agency's notice document URL; otherwise the workbook row itself.

    A tracking-sheet row that repeats an earlier row's employer, county,
    address, Date Received, projected date and employee count (``_reentry_key``)
    is a re-entry of that listing and is held as ``duplicate_row_in_source``
    with the columns that differ, unless the rows name different affected
    occupations (``_occupations_match``).
    """
    rows = read_archive(directory)
    existing_ids = {str(x) for x in existing_ids or ()}
    existing_urls = set(existing_urls or ())
    fields = [(row, _archive_fields(row)) for row in rows]
    # The 2025 CSV carries notice numbers, so its rows own their document URLs.
    numbered_urls = {f["document_url"]: row["source_row"] for row, f in fields
                     if row["file"] == REPORT_2025 and f["document_url"]}
    seen_urls: dict[str, str] = {}
    seen_ids: dict[str, str] = {}
    seen_content: dict[str, str] = {}
    seen_reentry: dict[str, list[dict]] = {}
    records, held = [], []
    for row, f in fields:
        url = f["document_url"]
        document = SALESFORCE_DOC.fullmatch(url or "")
        # Two identical cell sets on one tracking sheet are one listing.
        content_key = _json([row["sheet"], row["raw"]]) if row["file"] == TRACKING else None
        reentry_key = _reentry_key(row, f)
        reentry_of = next((other for other in seen_reentry.get(reentry_key, [])
                           if _occupations_match(other["raw"].get("Affected Occupations"),
                                                 row["raw"].get("Affected Occupations"))), None)
        text = " ".join([f["company"], f["comment"]])
        related = None
        differing = None
        if not f["company"]:
            reason = "missing_employer"
        elif OUT_OF_STATE_TEXT.search(f["county"] or ""):
            reason = "source_explicitly_out_of_state"
        elif row["file"] == REPORT_2025 and not f["notice_id"]:
            reason = "invalid_notice_number"
        elif f["notice_id"] and f["notice_id"] in existing_ids:
            reason = "already_represented_ky_id"
        elif f["notice_id"] and f["notice_id"] in seen_ids:
            reason, related = "duplicate_notice_number_in_source", seen_ids[f["notice_id"]]
        elif not f["notice_id"] and url and url in existing_urls:
            # A numbered row keeps its number as identity: the agency has
            # linked one document from two numbered notices (2657 and 2820).
            reason = "already_represented_ky_notice_url"
        elif row["file"] != REPORT_2025 and url and url in numbered_urls:
            reason, related = "listed_in_ky_2025_report_with_notice_number", numbered_urls[url]
        elif row["file"] != REPORT_2025 and document and url in seen_urls:
            reason, related = "duplicate_notice_url_in_source", seen_urls[url]
        elif f["source"] and "WARN" not in f["source"].upper():
            reason = "ky_tracking_source_not_warn"
        elif RESCINDED.search(text):
            reason = "agency_comment_rescinded"
        elif AMENDMENT.search(text):
            reason = "amendment_row_parent_unresolved"
        elif content_key is not None and content_key in seen_content:
            reason, related = "duplicate_row_in_source", seen_content[content_key]
        elif reentry_of is not None:
            reason, related = "duplicate_row_in_source", reentry_of["source_row"]
            differing = sorted(head for head in set(reentry_of["raw"]) | set(row["raw"])
                               if reentry_of["raw"].get(head) != row["raw"].get(head))
        else:
            reason = None
        if f["notice_id"]:
            seen_ids.setdefault(f["notice_id"], row["source_row"])
        if document and row["file"] != REPORT_2025:
            seen_urls.setdefault(url, row["source_row"])
        if content_key is not None:
            seen_content.setdefault(content_key, row["source_row"])
        if reentry_key is not None:
            seen_reentry.setdefault(reentry_key, []).append(row)
        if reason:
            item = {"origin": row["source_artifact"], "state": "KY", "reason": reason,
                    "source_row": row["source_row"],
                    "source_row_sha256": row["source_row_sha256"],
                    "source_notice_id": f["notice_id"], "source_url": row["source_url"],
                    "document_url": url, "notice_year": f["year"],
                    "raw_extra": _json({"cells": row["raw"],
                                        "stored_numbers": row["stored_numbers"]})}
            if related:
                item["related_source_row"] = related
            if differing is not None:
                item["differing_columns"] = differing
            held.append(item)
            continue
        if f["notice_id"]:
            identity, basis = f"KY:{f['notice_id']}", "agency_notice_number"
        elif document:
            identity, basis = f"KY:document:{document.group(1)}", "agency_notice_document_url"
        else:
            tag = "tracking" if row["file"] == TRACKING else "report"
            identity = f"KY:archive:{tag}:{row['sheet']}:r{row['row_number']}"
            basis = "agency_workbook_row"
        start, end = f["projected"]
        kind = f["kind"].casefold()
        details = {
            "source_artifact": row["source_artifact"],
            "source_artifact_sha256": row["source_artifact_sha256"],
            "source_row": row["source_row"], "source_row_sha256": row["source_row_sha256"],
            "source_document_url": url, "identity_basis": basis,
            "agency_received_date": f["received"],
            "projected_action_date": start, "projected_action_date_end": end,
            "date_roles": {"Date Received": "agency_receipt",
                           f["projected_field"]: "projected_action"},
            "date_evidence_rule": "ky_agency_archive_labeled_fields_v1",
            "location_basis": "agency_county_field" if f["county"] else None,
            "address_text": f["address"], "address_role": "unverified" if f["address"] else None,
            "raw_fields": row["raw"],
        }
        if row["stored_numbers"]:
            details["stored_cell_numbers"] = row["stored_numbers"]
            details["stored_cell_rule"] = "employees_cell_stored_number_displayed_as_date_v1"
        details = {k: v for k, v in details.items() if v is not None}
        record = {
            "state": "KY", "employer_name": f["company"], "location": f["county"],
            "notice_date": None,
            "effective_date": start, "effective_date_end": end,
            "employees_affected": f["workers"],
            "layoff_type": ("closure" if kind.startswith("closure") else
                            "mass_layoff" if kind.startswith("layoff") else "unknown"),
            "is_temporary": None, "is_amendment": 0,
            "source_url": row["source_url"], "source_notice_id": f["notice_id"],
            "source_identity": identity, "source_details": _json(details),
            "raw_extra": _json({"cells": row["raw"], "stored_numbers": row["stored_numbers"]}),
            "dedupe_key": hashlib.sha1(f"KY|source|{identity}".encode()).hexdigest(),
        }
        if start:
            record.update(effective_date_precision="day", effective_date_basis="reported")
        if end:
            record.update(effective_date_end_precision="day", effective_date_end_basis="reported")
        record["raw_record_hash"] = _record_hash(record)
        records.append(record)
    if len(records) + len(held) != len(rows) or len({r["source_identity"] for r in records}) != len(records):
        raise ValueError("Kentucky archive row accounting mismatch")
    by_file: dict[str, dict] = {}
    for row in rows:
        by_file.setdefault(row["file"], {"source_rows": 0})["source_rows"] += 1
    return records, held, {
        "source_rows": len(rows), "admitted": len(records), "held": len(held),
        "rows_by_file": {name: item["source_rows"] for name, item in sorted(by_file.items())},
        "hold_reasons": dict(sorted(Counter(item["reason"] for item in held).items())),
        "admitted_workers": sum(item["employees_affected"] or 0 for item in records),
        "admitted_missing_workers": sum(item["employees_affected"] is None for item in records),
        "admitted_by_identity_basis": dict(sorted(Counter(
            json.loads(item["source_details"])["identity_basis"] for item in records).items())),
    }
