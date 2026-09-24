"""Extract and conservatively project official Iowa WARN observations.

An Excel row is an evidence pointer, not necessarily an additive layoff. Some
rows describe phases; others amend dates or worker counts. A source row is
admitted only where its event identity is unambiguous.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from warnlive.normalize.engine import _fold, _record_hash
from warnlive.normalize.revisions import classify_idless_events

HEADERS = (
    "Company", "Street Address", "City", "County", "State", "ZIP",
    "Notice Type", "Number of Employees Affected", "Notice Date",
    "Layoff Date", "Local Workforce Area", "Industry",
)
PDF_COLUMN_STARTS = (0, 168, 254, 310, 353, 369, 399, 448, 467, 515, 564, 632, 792)
PDF_DATE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")


def _cell(value: object) -> dict:
    """Keep Excel value types in a JSON-safe, stable source projection."""
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        return {"type": "number", "value": value}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    if value is None:
        return {"type": "blank", "value": None}
    raise ValueError(f"unsupported Iowa cell value: {type(value).__name__}")


def _pdf_date(value: str) -> str | None:
    for pattern in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(value, pattern).date().isoformat()
        except ValueError:
            pass
    return None


def _is_amendment(value: str) -> bool:
    text = value.casefold()
    return any(marker in text for marker in ("amend", "additional employees", "change in"))


def extract(directory: Path) -> list[dict]:
    """Validate a pinned workbook, accounting for every populated data row."""
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if manifest.get("format") != "warn-ia-source-artifacts-v1":
        raise ValueError("unsupported Iowa source manifest")
    artifacts = manifest.get("artifacts") or []
    if {artifact.get("path") for artifact in artifacts} != {
        "event-log.xlsx", "historical-2023.pdf"
    } or len(artifacts) != 2:
        raise ValueError("Iowa source manifest must list both official logs")
    by_name = {artifact["path"]: artifact for artifact in artifacts}
    for artifact in artifacts:
        artifact_path = directory / artifact["path"]
        artifact_bytes = artifact_path.read_bytes()
        if (len(artifact_bytes) != artifact["bytes"] or
                hashlib.sha256(artifact_bytes).hexdigest() != artifact["sha256"]):
            raise ValueError(f"Iowa source checksum mismatch: {artifact_path}")
    historical = directory / "historical-2023.pdf"
    with pdfplumber.open(historical) as pdf:
        if len(pdf.pages) != by_name["historical-2023.pdf"]["pages"]:
            raise ValueError("Iowa historical PDF page count changed")
    item = by_name["event-log.xlsx"]
    path = directory / "event-log.xlsx"

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [item["sheet"]]:
            raise ValueError("Iowa workbook sheet layout changed")
        sheet = workbook[item["sheet"]]
        if sheet.max_column != len(HEADERS):
            raise ValueError("Iowa workbook column count changed")
        header_row = item["header_row"]
        header = tuple(cell.value for cell in sheet[header_row])
        if header != HEADERS:
            raise ValueError("Iowa workbook header changed")
        if "since June 2021" not in str(sheet["A4"].value):
            raise ValueError("Iowa workbook coverage statement changed")
        result = []
        trailing_blank = False
        for ordinal, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if all(cell is None for cell in cells):
                trailing_blank = True
                continue
            if trailing_blank:
                raise ValueError(f"Iowa data follows a blank row at {ordinal}")
            if len(cells) != len(HEADERS) or any(cell is None for cell in cells):
                raise ValueError(f"incomplete Iowa event-log row {ordinal}")
            company, street, city, county, state, zipcode, notice_type, workers, notice, layoff, area, industry = cells
            if not isinstance(workers, int) or isinstance(workers, bool) or workers < 0:
                raise ValueError(f"invalid Iowa worker count at row {ordinal}")
            if not isinstance(notice, (date, datetime)) or not isinstance(layoff, (date, datetime)):
                raise ValueError(f"invalid Iowa date at row {ordinal}")
            raw_cells = [_cell(value) for value in cells]
            raw = json.dumps(raw_cells, sort_keys=True, ensure_ascii=False,
                             separators=(",", ":"))
            result.append({
                "source_artifact": "agency/ia/event-log.xlsx",
                "source_row": f"event-log.xlsx:{item['sheet']}:r{ordinal}",
                "source_row_sha256": hashlib.sha256(raw.encode()).hexdigest(),
                "source_xlsx_sha256": item["sha256"],
                "source_url": item["source_url"],
                "raw_cells": raw_cells,
                "company_text": company,
                "street_address_text": street,
                "city_text": city,
                "county_text": county,
                "address_state_text": state,
                "postal_code_text": zipcode,
                "address_role": "unverified",
                "notice_type_text": notice_type,
                "workers_reported": workers,
                "notice_date": notice.date().isoformat() if isinstance(notice, datetime) else notice.isoformat(),
                "effective_date": layoff.date().isoformat() if isinstance(layoff, datetime) else layoff.isoformat(),
                "local_workforce_area_text": area,
                "industry_text": industry,
                "decision_status": "needs_correspondence_and_amendment_review",
            })
        if len(result) != item["data_rows"]:
            raise ValueError("Iowa workbook row count changed")
        return result
    finally:
        workbook.close()


def extract_historical(directory: Path) -> list[dict]:
    """Project each historical-PDF line with its original page and cells.

    The PDF is a fixed-column printed worksheet, not a structured table. The
    worker-count column anchors its rows, including a malformed notice date;
    unexpected layouts fail closed.
    """
    extract(directory)  # validates both artifacts against the pinned manifest
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    artifact = next(item for item in manifest["artifacts"]
                    if item["path"] == "historical-2023.pdf")
    result = []
    with pdfplumber.open(directory / "historical-2023.pdf") as pdf:
        for page_number, (page, expected) in enumerate(
            zip(pdf.pages, artifact["data_rows_by_page"], strict=True), start=1
        ):
            words = page.extract_words(x_tolerance=1, y_tolerance=3, use_text_flow=True)
            anchors = sorted(
                (word for word in words if 448 <= word["x0"] < 467
                 and word["text"].isdigit()),
                key=lambda word: word["top"],
            )
            if len(anchors) != expected:
                raise ValueError(f"Iowa historical PDF row count changed on page {page_number}")
            for ordinal, anchor in enumerate(anchors, start=1):
                # The worker column is present even where the printed notice
                # date is malformed. A long company can wrap onto preceding
                # lines, ending level with the other cells of this row.
                previous_top = anchors[ordinal - 2]["top"] if ordinal > 1 else None
                lower = previous_top + 2.5 if previous_top is not None else anchor["top"] - 4.5
                cells: list[list[str]] = [[] for _ in range(12)]
                for word in words:
                    if not lower <= word["top"] <= anchor["top"] + 1.5:
                        continue
                    for index, (left, right) in enumerate(
                        zip(PDF_COLUMN_STARTS, PDF_COLUMN_STARTS[1:])
                    ):
                        if left <= word["x0"] < right:
                            cells[index].append(word["text"])
                            break
                values = [" ".join(parts) for parts in cells]
                if not values[0] or not values[7] or not values[8]:
                    raise ValueError(f"incomplete Iowa historical row {page_number}:{ordinal}")
                if not values[7].isdigit():
                    raise ValueError(f"invalid Iowa historical workers {page_number}:{ordinal}")
                raw = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
                layout_issues = [f"empty_column_{index + 1}" for index, value in
                                 enumerate(values) if not value]
                notice_date, effective_date = _pdf_date(values[8]), _pdf_date(values[9])
                if notice_date is None:
                    layout_issues.append("invalid_notice_date")
                if effective_date is None:
                    layout_issues.append("invalid_layoff_date")
                result.append({
                    "source_artifact": "agency/ia/historical-2023.pdf",
                    "source_row": f"historical-2023.pdf:p{page_number}:r{ordinal}",
                    "source_row_sha256": hashlib.sha256(raw.encode()).hexdigest(),
                    "source_pdf_sha256": artifact["sha256"],
                    "source_url": artifact["source_url"],
                    "raw_cells": values,
                    "layout_issues": layout_issues,
                    "company_text": values[0], "street_address_text": values[1],
                    "city_text": values[2], "county_text": values[3],
                    "address_state_text": values[4], "postal_code_text": values[5],
                    "address_role": "unverified", "notice_type_text": values[6],
                    "workers_reported": int(values[7]),
                    "notice_date_text": values[8], "effective_date_text": values[9],
                    "notice_date": notice_date, "effective_date": effective_date,
                    "local_workforce_area_text": values[10],
                    "industry_text": values[11],
                    "decision_status": "needs_correspondence_and_amendment_review",
                })
    return result


def project(rows: list[dict]) -> tuple[list[dict], list[dict], dict, dict[str, str]]:
    """Admit uniquely identified ordinary rows; retain amendment and site questions.

    The two agency artifacts overlap. A printed PDF row and an Excel row with
    the same employer, dates, city and worker count are capture observations of
    one event, not two notices. Differing sites or worker allocations remain
    unresolved unless an agency filing identifier establishes their meaning.
    """
    if len({row["source_row"] for row in rows}) != len(rows):
        raise ValueError("duplicate Iowa source row pointer")
    families: dict[tuple[str, str | None, str | None], list[dict]] = defaultdict(list)
    for row in rows:
        families[(_fold(row["company_text"]), row["notice_date"],
                  row["effective_date"])].append(row)
    status: dict[str, str] = {}
    related: dict[str, str] = {}
    candidates: list[dict] = []
    for group in families.values():
        ordinary = [row for row in group if not _is_amendment(row["notice_type_text"])]
        for row in group:
            if _is_amendment(row["notice_type_text"]):
                status[row["source_row"]] = "amendment_without_verified_parent"
            elif row.get("layout_issues") or not row["notice_date"] or not row["effective_date"]:
                status[row["source_row"]] = "unresolved_source_layout_or_date"
        ordinary = [row for row in ordinary if row["source_row"] not in status]
        # Multiple worker totals or locations under one employer/date pair may
        # be sites, phases, or revisions. Do not sum or publish them separately.
        workers = {row["workers_reported"] for row in ordinary}
        cities = {_fold(row["city_text"]) for row in ordinary}
        addresses = {_fold(row["street_address_text"]) for row in ordinary}
        if len(workers) > 1 or len(cities) > 1 or len(addresses) > 1:
            for row in ordinary:
                status[row["source_row"]] = "site_phase_or_worker_allocation_unresolved"
            continue
        if ordinary:
            # Prefer the structured, current agency log when a historical PDF
            # prints the same event. Distinct rows in the same artifact remain
            # unresolved because matching totals do not prove one filing.
            by_artifact = Counter(row["source_artifact"] for row in ordinary)
            if any(count > 1 for count in by_artifact.values()):
                for row in ordinary:
                    status[row["source_row"]] = "repeated_source_event_unresolved"
                continue
            preferred = sorted(ordinary, key=lambda row:
                               (row["source_artifact"] != "agency/ia/event-log.xlsx",
                                row["source_row"]))[0]
            candidates.append(preferred)
            for row in ordinary:
                if row is not preferred:
                    status[row["source_row"]] = "duplicate_agency_capture"
                    related[row["source_row"]] = preferred["source_row"]

    decisions = classify_idless_events(
        candidates, row_key=lambda row: row["source_row"],
        employer=lambda row: _fold(row["company_text"]),
        site=lambda row: _fold(row["street_address_text"] or row["city_text"]),
        action=lambda row: row["effective_date"],
        notice=lambda row: row["notice_date"],
        content=lambda row: json.dumps((row["notice_date"], row["effective_date"],
                                        row["workers_reported"])),
    )
    for pointer, decision in decisions.items():
        status[pointer] = ("admitted" if decision.kind == "notice" else
                           "possible_revision_same_event" if decision.kind == "unresolved" else
                           "duplicate_agency_capture")
        if decision.related_row:
            related[pointer] = decision.related_row
    records, held = [], []
    for row in rows:
        pointer = row["source_row"]
        reason = status[pointer]
        row["disposition"] = reason
        if pointer in related:
            row["related_source_row"] = related[pointer]
        if reason != "admitted":
            held.append({"origin": row["source_artifact"], "state": "IA", "reason": reason,
                         "source_row": pointer, "source_row_sha256": row["source_row_sha256"],
                         "source_url": row["source_url"], "notice_year":
                         row["notice_date"][:4] if row["notice_date"] else None,
                         "related_source_row": related.get(pointer),
                         "raw_extra": json.dumps(row["raw_cells"], ensure_ascii=False, default=str)})
            continue
        identity = f"IA:agency-observation:{hashlib.sha256(pointer.encode()).hexdigest()}"
        details = {"source_artifact": row["source_artifact"], "source_row": pointer,
                   "source_row_sha256": row["source_row_sha256"],
                   "identity_basis": "unique_agency_event_observation_not_filing_id",
                   "notice_type_text": row["notice_type_text"],
                   "address_role": "unverified", "street_address_text": row["street_address_text"],
                   "city_text": row["city_text"], "county_text": row["county_text"],
                   "address_state_text": row["address_state_text"]}
        rec = {"state": "IA", "employer_name": row["company_text"].strip(),
               "location": None,
               "notice_date": row["notice_date"], "notice_date_precision": "day",
               "notice_date_basis": "reported", "effective_date": row["effective_date"],
               "effective_date_precision": "day", "effective_date_basis": "reported",
               "employees_affected": row["workers_reported"] or None,
               "layoff_type": "unknown", "is_temporary": None, "is_amendment": 0,
               "source_url": row["source_url"], "source_notice_id": pointer,
               "source_identity": identity, "source_details": json.dumps(details, sort_keys=True),
               "raw_extra": json.dumps(row["raw_cells"], ensure_ascii=False, default=str),
               "dedupe_key": hashlib.sha1(identity.encode()).hexdigest()}
        rec["raw_record_hash"] = _record_hash(rec)
        records.append(rec)
    if len(records) + len(held) != len(rows):
        raise ValueError("Iowa source row accounting mismatch")
    return records, held, {"source_rows": len(rows), "admitted": len(records),
                           "held": len(held), "hold_reasons": dict(Counter(x["reason"] for x in held))}, related
